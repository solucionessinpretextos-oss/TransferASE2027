from __future__ import annotations

import re
import textwrap
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
MANUALS = ROOT / "docs" / "manuales"
ASSETS = MANUALS / "assets"
OUTPUT = MANUALS / "entregables"
OUTPUT.mkdir(parents=True, exist_ok=True)

BROWN = "5B3112"
GOLD = "B98A36"
PALE_GOLD = "F5EEDC"
CHARCOAL = "191D24"
LIGHT = "F4F6F8"
MID = "D7DCE2"
GREEN = "0B8F67"
RED = "B83245"
BLUE = "176B87"
PURPLE = "6957A5"
ORANGE = "C46210"

FONT_REGULAR = Path(r"C:\Windows\Fonts\segoeui.ttf")
FONT_SEMIBOLD = Path(r"C:\Windows\Fonts\seguisb.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\segoeuib.ttf")


MANUAL_SPECS = [
    ("01-Manual-Usuario-TransferClienteASE.md", "Manual_Usuario_TransferClienteASE.docx", "Manual de usuario", "TransferClienteASE", "Personas usuarias y enlaces de entes fiscalizados"),
    ("02-Manual-Administrador-Dashboard.md", "Manual_Administrador_FileXfer_AdminDashboard.docx", "Manual de administración", "FileXfer.AdminDashboard", "Administradores y personal de operación"),
    ("03-Manual-Operacion-Servidor.md", "Manual_Operacion_FileXfer_Server_Full.docx", "Manual de operación", "FileXfer.Server.Full", "Infraestructura, soporte y responsables de despliegue"),
    ("04-Manual-Actualizador.md", "Manual_Actualizacion_TransferClienteASEUpdater.docx", "Manual del actualizador", "TransferClienteASEUpdater", "Soporte técnico y responsables de liberación"),
    ("05-Manual-Tecnico-Validacion-Archivos.md", "Manual_Tecnico_FileXfer_Pdf.docx", "Manual técnico", "FileXfer.Pdf", "Desarrollo, seguridad y operación"),
    ("06-Manual-Pruebas-Automatizadas.md", "Manual_Pruebas_Automatizadas_TransferASE.docx", "Manual de pruebas", "Proyectos SmokeTests", "Desarrollo, QA y responsables de liberación"),
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    source = FONT_BOLD if bold else FONT_REGULAR
    return ImageFont.truetype(str(source), size)


def rounded_box(draw: ImageDraw.ImageDraw, box, fill, outline=GOLD, radius=24, width=3):
    draw.rounded_rectangle(box, radius=radius, fill=f"#{fill}", outline=f"#{outline}", width=width)


def wrapped_center(draw: ImageDraw.ImageDraw, xy, text, max_chars, fill, size, bold=False, spacing=6):
    lines = textwrap.wrap(text, max_chars) or [text]
    f = font(size, bold)
    box = draw.multiline_textbbox((0, 0), "\n".join(lines), font=f, spacing=spacing, align="center")
    width = box[2] - box[0]
    height = box[3] - box[1]
    draw.multiline_text((xy[0] - width / 2, xy[1] - height / 2), "\n".join(lines), font=f, fill=f"#{fill}", spacing=spacing, align="center")


def arrow(draw: ImageDraw.ImageDraw, start, end, color=GOLD, width=8):
    draw.line([start, end], fill=f"#{color}", width=width)
    x2, y2 = end
    if abs(end[0] - start[0]) >= abs(end[1] - start[1]):
        sign = 1 if end[0] > start[0] else -1
        points = [(x2, y2), (x2 - sign * 22, y2 - 14), (x2 - sign * 22, y2 + 14)]
    else:
        sign = 1 if end[1] > start[1] else -1
        points = [(x2, y2), (x2 - 14, y2 - sign * 22), (x2 + 14, y2 - sign * 22)]
    draw.polygon(points, fill=f"#{color}")


def flow_diagram(path: Path, title: str, steps: list[tuple[str, str]], footer: str):
    img = Image.new("RGB", (1700, 820), "#F6F7F9")
    draw = ImageDraw.Draw(img)
    draw.text((80, 50), title, font=font(42, True), fill=f"#{CHARCOAL}")
    draw.text((82, 110), footer, font=font(22), fill="#5E6875")
    n = len(steps)
    gap = 34
    box_w = int((1540 - gap * (n - 1)) / n)
    y1, y2 = 255, 590
    for i, (head, detail) in enumerate(steps):
        x1 = 80 + i * (box_w + gap)
        x2 = x1 + box_w
        rounded_box(draw, (x1, y1, x2, y2), "FFFFFF", outline=GOLD)
        draw.ellipse((x1 + 24, y1 + 24, x1 + 84, y1 + 84), fill=f"#{BROWN}")
        wrapped_center(draw, (x1 + 54, y1 + 54), str(i + 1), 3, "FFFFFF", 26, True)
        wrapped_center(draw, ((x1 + x2) / 2, y1 + 145), head, 18, CHARCOAL, 28, True)
        wrapped_center(draw, ((x1 + x2) / 2, y1 + 245), detail, 25, "596575", 21)
        if i < n - 1:
            arrow(draw, (x2 + 7, (y1 + y2) / 2), (x2 + gap - 7, (y1 + y2) / 2))
    draw.rounded_rectangle((80, 670, 1620, 755), radius=18, fill=f"#{PALE_GOLD}")
    wrapped_center(draw, (850, 712), "Cada paso confirma el anterior. Si uno falla, no se simula éxito: se detiene, informa y conserva evidencia.", 110, BROWN, 24, True)
    img.save(path, optimize=True)


def architecture_diagram(path: Path):
    img = Image.new("RGB", (1700, 950), "#F6F7F9")
    draw = ImageDraw.Draw(img)
    draw.text((80, 50), "Arquitectura operativa de Transfer ASE", font=font(42, True), fill=f"#{CHARCOAL}")
    draw.text((82, 110), "Flujos autenticados, reanudables, verificables y auditables", font=font(23), fill="#5E6875")
    boxes = [
        ((80, 230, 430, 470), "Cliente WPF", "e.firma · catálogos\nverificación local\nprogreso y reanudación", BLUE),
        ((560, 230, 1140, 470), "Servidor FileXfer", "API REST + gRPC\nautorización · políticas\nrecibos · auditoría", BROWN),
        ((1270, 230, 1620, 470), "Dashboard", "monitoreo · KPI\npolíticas · trazabilidad\nroles admin/ops", PURPLE),
        ((370, 650, 760, 850), "FileXferDb", "sesiones, políticas, cargas,\nrecibos y outbox", GREEN),
        ((930, 650, 1320, 850), "BuDi_Dev", "auditorías, oficios,\nrequerimientos y catálogo", ORANGE),
        ((1310, 650, 1620, 850), "Storage", "parciales, finalizados\ny evidencia durable", BLUE),
    ]
    for box, head, detail, color in boxes:
        rounded_box(draw, box, "FFFFFF", outline=color)
        wrapped_center(draw, ((box[0] + box[2]) / 2, box[1] + 70), head, 25, color, 30, True)
        wrapped_center(draw, ((box[0] + box[2]) / 2, box[1] + 155), detail, 32, "596575", 22)
    arrow(draw, (430, 350), (550, 350), BLUE)
    arrow(draw, (1140, 350), (1260, 350), PURPLE)
    arrow(draw, (760, 740), (920, 740), GREEN)
    arrow(draw, (850, 470), (590, 640), GREEN)
    arrow(draw, (900, 470), (1120, 640), ORANGE)
    arrow(draw, (1040, 470), (1450, 640), BLUE)
    draw.text((80, 900), "La interfaz nunca es la fuente final de verdad: el servidor confirma bytes persistidos, catálogo y recibo.", font=font(22, True), fill=f"#{BROWN}")
    img.save(path, optimize=True)


def test_matrix_diagram(path: Path):
    img = Image.new("RGB", (1700, 940), "#F6F7F9")
    draw = ImageDraw.Draw(img)
    draw.text((80, 50), "Mapa de validación automatizada", font=font(42, True), fill=f"#{CHARCOAL}")
    draw.text((82, 110), "Cada ejecutable de pruebas protege un contrato crítico distinto", font=font(23), fill="#5E6875")
    items = [
        ("PDF", "válido · cifrado\nreparable · activo", BLUE),
        ("FileSafety", "formatos · contraseña\nmacros · estructura", GREEN),
        ("Security", "firma · hash · rutas\nrollback del updater", RED),
        ("Resume", "checkpoint · reinicio\nidempotencia", PURPLE),
        ("Upload E2E", "autenticación · carga\ncatálogo · recibo", ORANGE),
        ("Staging", "SQL · storage · TLS\nhealth · readiness", BROWN),
    ]
    for i, (head, detail, color) in enumerate(items):
        row, col = divmod(i, 3)
        x1 = 80 + col * 525
        y1 = 230 + row * 300
        x2, y2 = x1 + 450, y1 + 220
        rounded_box(draw, (x1, y1, x2, y2), "FFFFFF", outline=color)
        draw.rectangle((x1, y1, x1 + 18, y2), fill=f"#{color}")
        wrapped_center(draw, ((x1 + x2) / 2, y1 + 65), head, 20, color, 30, True)
        wrapped_center(draw, ((x1 + x2) / 2, y1 + 145), detail, 28, "596575", 22)
    draw.rounded_rectangle((80, 835, 1620, 905), radius=18, fill=f"#{PALE_GOLD}")
    wrapped_center(draw, (850, 870), "Regla de liberación: una falla bloquea el paquete; una prueba omitida debe quedar explícitamente justificada.", 112, BROWN, 23, True)
    img.save(path, optimize=True)


def annotate(path_in: Path, path_out: Path, points: list[tuple[int, int, int]]):
    image = Image.open(path_in).convert("RGB")
    draw = ImageDraw.Draw(image)
    scale = max(image.width / 1200, image.height / 900)
    radius = max(18, int(25 * scale))
    f = font(max(16, int(22 * scale)), True)
    for number, x, y in points:
        x = int(x * image.width)
        y = int(y * image.height)
        draw.ellipse((x - radius, y - radius, x + radius, y + radius), fill="#FFFFFF", outline=f"#{RED}", width=max(3, int(4 * scale)))
        label = str(number)
        box = draw.textbbox((0, 0), label, font=f)
        draw.text((x - (box[2] - box[0]) / 2, y - (box[3] - box[1]) / 2 - 2), label, font=f, fill=f"#{RED}")
    image.save(path_out, optimize=True)


def create_visual_assets():
    diagram_dir = ASSETS / "diagramas"
    diagram_dir.mkdir(parents=True, exist_ok=True)
    architecture_diagram(diagram_dir / "arquitectura-ecosistema.png")
    flow_diagram(
        diagram_dir / "flujo-carga.png",
        "Flujo seguro de una carga",
        [("Seleccionar", "Archivo o carpeta y requerimiento correcto"), ("Inspeccionar", "Estructura, contraseña, contenido activo y política"), ("Transferir", "Bloques reanudables con progreso visible"), ("Confirmar", "Hash, catálogo y recibo digital")],
        "La aplicación informa la acción actual y el porcentaje de cada archivo y del lote.",
    )
    flow_diagram(
        diagram_dir / "flujo-actualizacion.png",
        "Actualización segura del cliente",
        [("Consultar", "Versión, obligatoriedad y notas"), ("Descargar", "Paquete desde HTTPS"), ("Verificar", "SHA-256 y firma RSA-PSS"), ("Aplicar", "Staging, respaldo, copia y rollback")],
        "El actualizador no confía en el nombre del ZIP; verifica su contenido antes de tocar la instalación.",
    )
    flow_diagram(
        diagram_dir / "flujo-validacion.png",
        "Decisión de seguridad de archivos",
        [("Política", "Extensión habilitada y límite opcional"), ("Inspección", "Formato real, estructura, cifrado y contenido activo"), ("Reparación", "Solo PDF reparable, siempre sobre copia"), ("Revalidación", "El resultado reparado debe aprobar toda la inspección")],
        "Cifrado, daño no reparable o estructura ambigua se rechazan; no existe el estado “más o menos válido”.",
    )
    test_matrix_diagram(diagram_dir / "mapa-pruebas.png")
    annotate(
        ASSETS / "cliente" / "01-login.png",
        ASSETS / "cliente" / "01-login-anotado.png",
        [(1, .25, .52), (2, .25, .605), (3, .25, .685), (4, .25, .773), (5, .49, .85)],
    )
    annotate(
        ASSETS / "cliente" / "05-carga-completada.png",
        ASSETS / "cliente" / "05-carga-completada-anotada.png",
        [(1, .02, .31), (2, .51, .31), (3, .67, .42), (4, .83, .42), (5, .83, .97)],
    )
    annotate(
        ASSETS / "dashboard" / "01-dashboard-general.png",
        ASSETS / "dashboard" / "01-dashboard-general-anotado.png",
        [(1, .08, .21), (2, .62, .21), (3, .31, .47), (4, .47, .435)],
    )
    annotate(
        ASSETS / "dashboard" / "06-politicas-archivos.png",
        ASSETS / "dashboard" / "06-politicas-archivos-anotada.png",
        [(1, .91, .055), (2, .82, .28), (3, .48, .35), (4, .79, .35)],
    )


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=MID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    """Evita que Word parta una fila entre dos páginas."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = OxmlElement("w:keepNext")
    p_pr.append(keep_next)


def remove_paragraph_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def set_image_alt(inline_shape, alt: str):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", alt)
    doc_pr.set("title", alt[:120])


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def configure_document(doc: Document, short_title: str):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(.62)
    section.bottom_margin = Inches(.62)
    section.left_margin = Inches(.72)
    section.right_margin = Inches(.72)
    section.header_distance = Inches(.25)
    section.footer_distance = Inches(.25)
    section.different_first_page_header_footer = False
    # Forzamos el valor falso en el XML, no solo la ausencia del elemento.
    # Algunas instalaciones de Word pueden heredar "pares e impares diferentes"
    # de Normal.dotm cuando el valor no está escrito explícitamente.
    doc.settings.odd_and_even_pages_header_footer = False
    settings_element = doc.settings.element
    even_odd = settings_element.find(qn("w:evenAndOddHeaders"))
    if even_odd is None:
        even_odd = OxmlElement("w:evenAndOddHeaders")
        settings_element.append(even_odd)
    even_odd.set(qn("w:val"), "false")

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(CHARCOAL)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size in (("Title", 30), ("Subtitle", 14), ("Heading 1", 21), ("Heading 2", 15), ("Heading 3", 12)):
        style = styles[style_name]
        style.font.name = "Segoe UI"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(size)
        if style_name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.keep_with_next = True
        p_pr = style.element.get_or_add_pPr()
        borders = p_pr.find(qn("w:pBdr"))
        if borders is not None:
            p_pr.remove(borders)

    def populate_header(header):
        header.paragraphs[0].text = ""
        table = header.add_table(rows=1, cols=2, width=Inches(7.05))
        table.columns[0].width = Inches(1.35)
        table.columns[1].width = Inches(5.7)
        header_logo = table.cell(0, 0).paragraphs[0].add_run().add_picture(
            str(ASSETS / "logo-ase.png"), width=Inches(1.25)
        )
        set_image_alt(header_logo, "Logotipo de la Auditoría Superior del Estado de Sinaloa")
        p = table.cell(0, 1).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(f"Transfer ASE 2027  ·  {short_title}")
        r.font.name = "Segoe UI"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(BROWN)

    def populate_footer(footer):
        footer.paragraphs[0].text = ""
        footer_table = footer.add_table(rows=1, cols=2, width=Inches(7.05))
        left = footer_table.cell(0, 0).paragraphs[0]
        left.add_run("Uso institucional  ·  Documentación controlada").font.size = Pt(8)
        right = footer_table.cell(0, 1).paragraphs[0]
        add_page_number(right)

    populate_header(section.header)
    populate_footer(section.footer)


def add_cover(doc: Document, category: str, component: str, audience: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(ASSETS / "logo-ase.png"), width=Inches(4.9))
    set_image_alt(shape, "Logotipo de la Auditoría Superior del Estado de Sinaloa")
    doc.add_paragraph("")
    title = doc.add_paragraph(category, style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_border(title)
    subtitle = doc.add_paragraph(component, style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.bold = True
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(BROWN)
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = band.cell(0, 0)
    set_cell_shading(cell, PALE_GOLD)
    set_cell_border(cell, GOLD, "10")
    pp = cell.paragraphs[0]
    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = pp.add_run("Guía gráfica, detallada y operable paso a paso")
    rr.bold = True
    rr.font.size = Pt(13)
    rr.font.color.rgb = RGBColor.from_string(BROWN)
    doc.add_paragraph("")
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields = [
        ("Audiencia", audience),
        ("Versión documental", "1.0"),
        ("Fecha", "22 de octubre de 2026"),
        ("Clasificación", "Uso institucional"),
    ]
    for row, (label, value) in zip(meta.rows, fields):
        row.cells[0].text = label
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].bold = True
        set_cell_shading(row.cells[0], LIGHT)
        for c in row.cells:
            set_cell_border(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run("Regla de oro: no avance si la pantalla no confirma el paso anterior.")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(RED)
    doc.add_page_break()


def add_contents(doc: Document, md_text: str):
    doc.add_heading("Contenido", level=1)
    doc.add_paragraph("Este índice sigue el orden operativo del manual. Los títulos también forman la navegación del documento.")
    heading_number = 0
    for line in md_text.splitlines():
        if line.startswith("## "):
            heading_number += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.25)
            p.paragraph_format.first_line_indent = Inches(-.25)
            p.add_run(f"{heading_number}. ").bold = True
            r = p.add_run(line[3:].strip())
            r.bold = True
        elif line.startswith("### "):
            p = doc.add_paragraph(style="List Bullet 2")
            p.add_run(line[4:].strip())
    doc.add_page_break()


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`.+?`|\[.+?\]\(.+?\))")


def add_inline(paragraph, text: str):
    pos = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor.from_string(BROWN)
        else:
            label, url = re.match(r"\[(.+?)\]\((.+?)\)", token).groups()
            run = paragraph.add_run(f"{label} ({url})")
            run.font.color.rgb = RGBColor.from_string(BLUE)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_callout(doc: Document, lines: list[str]):
    text = " ".join(line.lstrip("> ").strip() for line in lines)
    table = doc.add_table(rows=1, cols=1)
    set_row_cant_split(table.rows[0])
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    tone = RED if any(k in text.upper() for k in ("PELIGRO", "NO ", "BLOQUE", "ERROR")) else GOLD
    set_cell_shading(cell, "FFF8E8" if tone == GOLD else "FDEDEF")
    set_cell_border(cell, tone, "12")
    p = cell.paragraphs[0]
    add_inline(p, text)
    p.paragraph_format.space_after = Pt(0)


def add_code_block(doc: Document, lines: list[str]):
    table = doc.add_table(rows=1, cols=1)
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F0F2F4")
    set_cell_border(cell, "B9C0C8", "6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)


def add_markdown_table(doc: Document, lines: list[str]):
    rows = [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines]
    header = rows[0]
    data = rows[2:]
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for idx, text in enumerate(header):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, BROWN)
        set_cell_border(cell, BROWN)
        p = cell.paragraphs[0]
        add_inline(p, text)
        for r in p.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(8.7)
    for row_idx, values in enumerate(data):
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for idx, text in enumerate(values):
            set_cell_border(cells[idx])
            if row_idx % 2:
                set_cell_shading(cells[idx], "F8F9FA")
            p = cells[idx].paragraphs[0]
            add_inline(p, text)
            for r in p.runs:
                r.font.size = Pt(8.4)


def add_picture(doc: Document, md_file: Path, alt: str, raw_path: str):
    image_path = (md_file.parent / raw_path).resolve()
    if not image_path.exists():
        raise FileNotFoundError(f"No existe la imagen {image_path}")
    with Image.open(image_path) as img:
        ratio = img.height / img.width
    # El diagrama introductorio comparte página con una tabla operativa.
    # Mantenerlo ligeramente más compacto evita dejar una sola fila huérfana
    # en la página siguiente sin sacrificar legibilidad.
    width = 6.2 if Path(raw_path).name == "flujo-carga.png" else 6.95
    if ratio * width > 7.3:
        width = 7.3 / ratio
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(image_path), width=Inches(width))
    set_image_alt(shape, alt)


def parse_markdown(doc: Document, md_file: Path, text: str):
    lines = text.splitlines()
    index = 0
    first_h1_skipped = False
    code_lines = None
    heading_one_seen = False
    suppress_next_section_break = False
    while index < len(lines):
        line = lines[index].rstrip()
        if code_lines is not None:
            if line.startswith("```"):
                add_code_block(doc, code_lines)
                code_lines = None
            else:
                code_lines.append(line)
            index += 1
            continue
        if line.startswith("```"):
            code_lines = []
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if line.strip() == "<!-- continue-page -->":
            suppress_next_section_break = True
            index += 1
            continue
        if line.startswith("# ") and not first_h1_skipped:
            first_h1_skipped = True
            index += 1
            continue
        if line.startswith("## "):
            if heading_one_seen and not suppress_next_section_break:
                doc.add_page_break()
            suppress_next_section_break = False
            heading_one_seen = True
            p = doc.add_heading(line[3:].strip(), level=1)
            set_keep_with_next(p)
            index += 1
            continue
        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=2)
            set_keep_with_next(p)
            index += 1
            continue
        if line.startswith("#### "):
            p = doc.add_heading(line[5:].strip(), level=3)
            set_keep_with_next(p)
            index += 1
            continue
        image_match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if image_match:
            add_picture(doc, md_file, image_match.group(1), image_match.group(2))
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[index + 1]):
            table_lines = [line]
            index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].rstrip())
                index += 1
            add_markdown_table(doc, table_lines)
            continue
        if line.startswith(">"):
            quote_lines = [line]
            index += 1
            while index < len(lines) and lines[index].startswith(">"):
                quote_lines.append(lines[index])
                index += 1
            add_callout(doc, quote_lines)
            continue
        number_match = re.match(r"^(\d+)\.\s+(.*)", line)
        if number_match:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.32)
            p.paragraph_format.first_line_indent = Inches(-.28)
            p.add_run(f"{number_match.group(1)}. ").bold = True
            add_inline(p, number_match.group(2))
            index += 1
            continue
        bullet_match = re.match(r"^\s*-\s+(.*)", line)
        if bullet_match:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, bullet_match.group(1))
            index += 1
            continue
        if line.startswith("_") and line.endswith("_"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line.strip("_"))
            r.italic = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor.from_string("657180")
            index += 1
            continue
        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            nxt = lines[index].rstrip()
            if not nxt.strip() or nxt.startswith(("#", ">", "```", "|", "![")) or re.match(r"^(\d+)\.\s+", nxt) or re.match(r"^\s*-\s+", nxt):
                break
            paragraph_lines.append(nxt)
            index += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(x.strip() for x in paragraph_lines))
    if code_lines is not None:
        add_code_block(doc, code_lines)


def build_manual(md_name: str, docx_name: str, category: str, component: str, audience: str):
    md_file = MANUALS / md_name
    text = md_file.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc, component)
    add_cover(doc, category, component, audience)
    add_contents(doc, text)
    parse_markdown(doc, md_file, text)
    properties = doc.core_properties
    properties.title = f"{category} — {component}"
    properties.subject = "Ecosistema Transfer ASE 2027"
    properties.author = "Auditoría Superior del Estado de Sinaloa"
    properties.comments = "Documento generado desde fuentes versionadas del repositorio TransferASE2027."
    output = OUTPUT / docx_name
    doc.save(output)
    return output


def main():
    create_visual_assets()
    outputs = [build_manual(*spec) for spec in MANUAL_SPECS]
    for path in outputs:
        print(path)


if __name__ == "__main__":
    main()
